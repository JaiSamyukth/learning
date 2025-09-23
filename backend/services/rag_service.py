"""
RAG Service for Document Processing and Question Answering
Production-grade service for intelligent document chunking, embedding, and retrieval
"""

import asyncio
import re
import uuid
import json
import logging
from typing import List, Dict, Any, Optional, Tuple
from datetime import datetime, timedelta
from pathlib import Path

import nltk
from nltk.tokenize import sent_tokenize
from sentence_transformers import SentenceTransformer
import fitz  # PyMuPDF for PDF processing
import docx
import ebooklib
from ebooklib import epub
from bs4 import BeautifulSoup

from services.vector_service import vector_service
from services.database_service import database_service
from config.qdrant_config import qdrant_config
from utils.logger import get_logger


class RAGProcessingError(Exception):
    """Custom exception for RAG processing errors"""
    pass


class DocumentProcessor:
    """Document processing utilities for various file formats"""

    def __init__(self):
        self.logger = get_logger(__name__ + ".DocumentProcessor")

    async def extract_text(self, file_path: str, filename: str) -> Dict[str, Any]:
        """Extract text from various document formats"""
        try:
            file_extension = Path(filename).suffix.lower()

            if file_extension == '.pdf':
                return await self._extract_from_pdf(file_path)
            elif file_extension in ['.docx', '.doc']:
                return await self._extract_from_docx(file_path)
            elif file_extension == '.txt':
                return await self._extract_from_txt(file_path)
            elif file_extension == '.epub':
                return await self._extract_from_epub(file_path)
            elif file_extension in ['.html', '.htm']:
                return await self._extract_from_html(file_path)
            else:
                raise RAGProcessingError(f"Unsupported file format: {file_extension}")

        except Exception as e:
            self.logger.error(f"Text extraction failed for {filename}: {str(e)}")
            raise RAGProcessingError(f"Text extraction failed: {str(e)}")

    async def _extract_from_pdf(self, file_path: str) -> Dict[str, Any]:
        """Extract text from PDF with metadata"""
        doc = fitz.open(file_path)
        text_content = []

        for page_num in range(len(doc)):
            page = doc.load_page(page_num)
            text = page.get_text()
            text_content.append({
                'page': page_num + 1,
                'content': text,
                'word_count': len(text.split()) if text else 0
            })

        doc.close()

        total_text = '\n'.join([page['content'] for page in text_content])
        total_words = sum([page['word_count'] for page in text_content])

        return {
            'content': total_text,
            'pages': len(text_content),
            'total_words': total_words,
            'total_chars': len(total_text),
            'pages_content': text_content
        }

    async def _extract_from_docx(self, file_path: str) -> Dict[str, Any]:
        """Extract text from DOCX files"""
        doc = docx.Document(file_path)
        text_content = []

        for para in doc.paragraphs:
            if para.text.strip():
                text_content.append(para.text)

        total_text = '\n'.join(text_content)
        total_words = len(total_text.split())

        return {
            'content': total_text,
            'pages': len(doc.paragraphs) // 30,  # Rough page estimation
            'total_words': total_words,
            'total_chars': len(total_text),
            'pages_content': [{'page': 1, 'content': total_text, 'word_count': total_words}]
        }

    async def _extract_from_txt(self, file_path: str) -> Dict[str, Any]:
        """Extract text from plain text files"""
        with open(file_path, 'r', encoding='utf-8') as f:
            content = f.read()

        words = content.split()
        estimated_pages = len(words) // 300  # Rough estimation

        return {
            'content': content,
            'pages': max(1, estimated_pages),
            'total_words': len(words),
            'total_chars': len(content),
            'pages_content': [{'page': 1, 'content': content, 'word_count': len(words)}]
        }

    async def _extract_from_epub(self, file_path: str) -> Dict[str, Any]:
        """Extract text from EPUB files"""
        book = epub.read_epub(file_path)
        text_content = []

        for item in book.get_items():
            if item.get_type() == ebooklib.ITEM_DOCUMENT:
                content = item.get_content().decode('utf-8')
                soup = BeautifulSoup(content, 'html.parser')
                text = soup.get_text()
                if text.strip():
                    text_content.append(text)

        total_text = '\n'.join(text_content)
        total_words = len(total_text.split())

        return {
            'content': total_text,
            'pages': len(text_content),
            'total_words': total_words,
            'total_chars': len(total_text),
            'pages_content': [{'page': i+1, 'content': text, 'word_count': len(text.split())}
                            for i, text in enumerate(text_content)]
        }

    async def _extract_from_html(self, file_path: str) -> Dict[str, Any]:
        """Extract text from HTML files"""
        with open(file_path, 'r', encoding='utf-8') as f:
            content = f.read()

        soup = BeautifulSoup(content, 'html.parser')
        text = soup.get_text()
        words = text.split()
        estimated_pages = len(words) // 300

        return {
            'content': text,
            'pages': max(1, estimated_pages),
            'total_words': len(words),
            'total_chars': len(text),
            'pages_content': [{'page': 1, 'content': text, 'word_count': len(words)}]
        }


class TextChunker:
    """Intelligent text chunking for RAG"""

    def __init__(self):
        self.logger = get_logger(__name__ + ".TextChunker")

    def chunk_document(self, text: str, config: Dict[str, Any]) -> List[Dict[str, Any]]:
        """Chunk document using multiple strategies"""
        chunks = []

        # Strategy 1: Section-based chunking
        section_chunks = self._chunk_by_sections(text)
        chunks.extend(section_chunks)

        # Strategy 2: Paragraph-based chunking
        paragraph_chunks = self._chunk_by_paragraphs(text, config)
        chunks.extend(paragraph_chunks)

        # Strategy 3: Sentence-based chunking for dense content
        sentence_chunks = self._chunk_by_sentences(text, config)
        chunks.extend(sentence_chunks)

        # Remove duplicates and sort by position
        unique_chunks = self._remove_duplicate_chunks(chunks)

        # Add importance scoring
        scored_chunks = self._add_importance_scores(unique_chunks)

        return scored_chunks

    def _chunk_by_sections(self, text: str) -> List[Dict[str, Any]]:
        """Chunk by document sections (chapters, headings)"""
        chunks = []

        # Pattern for detecting sections
        section_patterns = [
            r'^Chapter\s+\d+.*$',
            r'^Section\s+\d+.*$',
            r'^[A-Z][^.!?]*[:.]',
            r'^\d+\.\s+[A-Z]',
            r'^#{1,6}\s+.*$'
        ]

        lines = text.split('\n')
        current_section = []
        current_section_title = "Introduction"

        for line in lines:
            line = line.strip()
            if not line:
                continue

            # Check if line is a section header
            is_section_header = any(re.match(pattern, line, re.MULTILINE | re.IGNORECASE)
                                  for pattern in section_patterns)

            if is_section_header and current_section:
                # Save previous section
                section_content = '\n'.join(current_section)
                if section_content.strip():
                    chunks.append({
                        'content': section_content,
                        'chunk_type': 'section',
                        'title': current_section_title,
                        'start_position': text.find(section_content),
                        'end_position': text.find(section_content) + len(section_content)
                    })

                current_section = []
                current_section_title = line

            current_section.append(line)

        # Add final section
        if current_section:
            section_content = '\n'.join(current_section)
            if section_content.strip():
                chunks.append({
                    'content': section_content,
                    'chunk_type': 'section',
                    'title': current_section_title,
                    'start_position': text.find(section_content),
                    'end_position': text.find(section_content) + len(section_content)
                })

        return chunks

    def _chunk_by_paragraphs(self, text: str, config: Dict[str, Any]) -> List[Dict[str, Any]]:
        """Chunk by paragraphs with smart overlap"""
        chunks = []

        paragraphs = [p.strip() for p in text.split('\n\n') if p.strip()]
        chunk_size = config.get('chunk_size', qdrant_config.DEFAULT_CHUNK_SIZE)
        chunk_overlap = config.get('chunk_overlap', qdrant_config.DEFAULT_CHUNK_OVERLAP)

        current_chunk = []
        current_length = 0

        for i, paragraph in enumerate(paragraphs):
            paragraph_length = len(paragraph)

            if current_length + paragraph_length > chunk_size and current_chunk:
                # Save current chunk
                chunk_content = '\n\n'.join(current_chunk)
                chunks.append({
                    'content': chunk_content,
                    'chunk_type': 'paragraph',
                    'paragraph_count': len(current_chunk),
                    'start_position': text.find(chunk_content),
                    'end_position': text.find(chunk_content) + len(chunk_content)
                })

                # Start new chunk with overlap
                overlap_start = max(0, len(current_chunk) - chunk_overlap)
                current_chunk = current_chunk[overlap_start:]
                current_length = sum(len(p) for p in current_chunk)

            current_chunk.append(paragraph)
            current_length += paragraph_length

        # Add final chunk
        if current_chunk:
            chunk_content = '\n\n'.join(current_chunk)
            chunks.append({
                'content': chunk_content,
                'chunk_type': 'paragraph',
                'paragraph_count': len(current_chunk),
                'start_position': text.find(chunk_content),
                'end_position': text.find(chunk_content) + len(chunk_content)
            })

        return chunks

    def _chunk_by_sentences(self, text: str, config: Dict[str, Any]) -> List[Dict[str, Any]]:
        """Chunk by sentences for dense content"""
        try:
            sentences = sent_tokenize(text)
        except:
            # Fallback if NLTK tokenization fails
            sentences = [s.strip() for s in text.split('.') if s.strip()]

        chunks = []
        chunk_size = config.get('chunk_size', qdrant_config.DEFAULT_CHUNK_SIZE)
        chunk_overlap = config.get('chunk_overlap', qdrant_config.DEFAULT_CHUNK_OVERLAP)

        current_chunk = []
        current_length = 0

        for i, sentence in enumerate(sentences):
            sentence_length = len(sentence)

            if current_length + sentence_length > chunk_size and current_chunk:
                # Save current chunk
                chunk_content = '. '.join(current_chunk)
                if not chunk_content.endswith('.'):
                    chunk_content += '.'

                chunks.append({
                    'content': chunk_content,
                    'chunk_type': 'sentence',
                    'sentence_count': len(current_chunk),
                    'start_position': text.find(chunk_content),
                    'end_position': text.find(chunk_content) + len(chunk_content)
                })

                # Start new chunk with overlap
                overlap_start = max(0, len(current_chunk) - chunk_overlap)
                current_chunk = current_chunk[overlap_start:]
                current_length = sum(len(s) for s in current_chunk)

            current_chunk.append(sentence)
            current_length += sentence_length

        # Add final chunk
        if current_chunk:
            chunk_content = '. '.join(current_chunk)
            if not chunk_content.endswith('.'):
                chunk_content += '.'

            chunks.append({
                'content': chunk_content,
                'chunk_type': 'sentence',
                'sentence_count': len(current_chunk),
                'start_position': text.find(chunk_content),
                'end_position': text.find(chunk_content) + len(chunk_content)
            })

        return chunks

    def _remove_duplicate_chunks(self, chunks: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        """Remove duplicate chunks and sort by position"""
        seen_content = set()
        unique_chunks = []

        for chunk in chunks:
            content_hash = hash(chunk['content'])
            if content_hash not in seen_content:
                seen_content.add(content_hash)
                unique_chunks.append(chunk)

        # Sort by start position
        unique_chunks.sort(key=lambda x: x.get('start_position', 0))
        return unique_chunks

    def _add_importance_scores(self, chunks: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        """Add importance scores to chunks based on content characteristics"""
        for chunk in chunks:
            content = chunk['content'].lower()
            score = 0.5  # Base score

            # Boost score for chunks with important keywords
            important_indicators = [
                'introduction', 'conclusion', 'summary', 'key points',
                'important', 'note', 'definition', 'theorem', 'principle'
            ]

            for indicator in important_indicators:
                if indicator in content:
                    score += 0.1

            # Boost score for section headers
            if chunk['chunk_type'] == 'section':
                score += 0.2

            # Boost score for longer, more substantial chunks
            if len(chunk['content']) > qdrant_config.DEFAULT_CHUNK_SIZE:
                score += 0.1

            # Normalize score between 0 and 1
            chunk['importance_score'] = min(1.0, max(0.1, score))

        return chunks


class RAGService:
    """Main RAG service for document processing and Q&A"""

    def __init__(self):
        self.logger = get_logger(__name__)
        self.document_processor = DocumentProcessor()
        self.text_chunker = TextChunker()

    async def process_document(self, file_path: str, filename: str, user_id: str) -> Dict[str, Any]:
        """Complete document processing pipeline"""
        try:
            # Step 1: Extract text
            self.logger.info(f"Starting document processing for {filename}")
            extraction_result = await self.document_processor.extract_text(file_path, filename)

            # Step 2: Get user RAG configuration
            rag_config = database_service.get_rag_config(int(user_id))
            if not rag_config:
                rag_config = {
                    'chunk_size': qdrant_config.DEFAULT_CHUNK_SIZE,
                    'chunk_overlap': qdrant_config.DEFAULT_CHUNK_OVERLAP
                }

            # Step 3: Create document record in database
            document_id = database_service.create_document(
                user_id=int(user_id),
                filename=filename,
                file_size=Path(file_path).stat().st_size,
                total_pages=extraction_result['pages'],
                total_tokens=extraction_result['total_words'],
                metadata={
                    'extraction_method': 'rag_processor',
                    'processing_timestamp': datetime.now().isoformat(),
                    'content_stats': {
                        'total_chars': extraction_result['total_chars'],
                        'total_words': extraction_result['total_words']
                    }
                }
            )

            # Step 4: Chunk the document
            chunks = self.text_chunker.chunk_document(extraction_result['content'], rag_config)

            # Step 5: Prepare chunks for database
            processed_chunks = []
            for i, chunk in enumerate(chunks):
                chunk_id = str(uuid.uuid4())

                # Calculate token count (approximate)
                token_count = len(chunk['content'].split())

                db_chunk = {
                    'id': chunk_id,
                    'document_id': document_id,
                    'user_id': int(user_id),
                    'content': chunk['content'],
                    'chunk_type': chunk['chunk_type'],
                    'importance_score': chunk['importance_score'],
                    'token_count': token_count,
                    'start_position': chunk.get('start_position', 0),
                    'end_position': chunk.get('end_position', len(chunk['content'])),
                    'metadata': {
                        'title': chunk.get('title', ''),
                        'paragraph_count': chunk.get('paragraph_count', 1),
                        'sentence_count': chunk.get('sentence_count', 1),
                        'processing_timestamp': datetime.now().isoformat()
                    }
                }

                database_service.create_chunk(**db_chunk)
                processed_chunks.append(db_chunk)

            # Step 6: Generate embeddings and store in Qdrant
            vector_stats = await vector_service.process_document_chunks(processed_chunks, user_id)

            # Step 7: Update document processing status
            database_service.execute_query('''
                UPDATE documents
                SET rag_processed = TRUE, updated_at = CURRENT_TIMESTAMP
                WHERE id = ?
            ''', (document_id,))

            # Step 8: Log processing completion
            database_service.log_system_metric(
                'document_processed',
                1,
                {
                    'document_id': document_id,
                    'user_id': user_id,
                    'filename': filename,
                    'total_chunks': len(processed_chunks),
                    'total_tokens': sum(c['token_count'] for c in processed_chunks)
                }
            )

            result = {
                'document_id': document_id,
                'filename': filename,
                'total_chunks': len(processed_chunks),
                'total_tokens': sum(c['token_count'] for c in processed_chunks),
                'processing_time': datetime.now().isoformat(),
                'vector_stats': vector_stats,
                'status': 'completed'
            }

            self.logger.info(f"Document processing completed successfully: {filename}")
            return result

        except Exception as e:
            self.logger.error(f"Document processing failed for {filename}: {str(e)}")
            raise RAGProcessingError(f"Document processing failed: {str(e)}")

    async def generate_answer(self, query: str, user_id: str, context_docs: List[str] = None) -> Dict[str, Any]:
        """Generate answer using RAG"""
        try:
            start_time = datetime.now()

            # Step 1: Retrieve relevant chunks
            relevant_chunks = await vector_service.search_similar(query, user_id)

            if not relevant_chunks:
                return {
                    'answer': "I don't have enough context to answer this question. Please upload a relevant document first.",
                    'sources': [],
                    'confidence': 0.0,
                    'chunks_used': 0
                }

            # Step 2: Assemble context
            context = self._assemble_context(relevant_chunks)

            # Step 3: Generate answer using chat service (with RAG context)
            from backend.services.chat_service import chat_service
            response = await chat_service.generate_response(query, context, user_id)

            # Step 4: Calculate response time
            response_time = (datetime.now() - start_time).total_seconds()

            # Step 5: Log interaction
            chunk_ids = [chunk['id'] for chunk in relevant_chunks[:5]]  # Top 5 chunks
            database_service.log_user_interaction(
                user_id=int(user_id),
                interaction_type='rag_chat',
                query=query,
                response=response['response'],
                rag_chunks_used=chunk_ids,
                performance_score=min(1.0, len(relevant_chunks) / 5.0),  # Simple scoring
                response_time=response_time,
                tokens_used=response.get('tokens_used', 0),
                metadata={
                    'chunk_count': len(relevant_chunks),
                    'avg_chunk_score': sum(c['score'] for c in relevant_chunks) / len(relevant_chunks),
                    'context_length': len(context)
                }
            )

            return {
                'answer': response['response'],
                'sources': chunk_ids,
                'confidence': min(1.0, len(relevant_chunks) / 5.0),
                'chunks_used': len(relevant_chunks),
                'context_length': len(context),
                'response_time': response_time
            }

        except Exception as e:
            self.logger.error(f"RAG answer generation failed: {str(e)}")
            raise RAGProcessingError(f"Answer generation failed: {str(e)}")

    def _assemble_context(self, relevant_chunks: List[Dict[str, Any]], max_context_length: int = 4000) -> str:
        """Assemble context from relevant chunks"""
        context_parts = []
        current_length = 0

        # Sort chunks by importance score and relevance
        sorted_chunks = sorted(relevant_chunks, key=lambda x: (x['score'], x.get('importance_score', 0.5)), reverse=True)

        for chunk in sorted_chunks:
            chunk_content = chunk['content']
            chunk_length = len(chunk_content)

            if current_length + chunk_length > max_context_length:
                # Truncate chunk if necessary
                remaining_length = max_context_length - current_length
                if remaining_length > 100:  # Only add if meaningful content remains
                    context_parts.append(chunk_content[:remaining_length])
                break

            context_parts.append(chunk_content)
            current_length += chunk_length

        return '\n\n'.join(context_parts)

    async def get_document_chunks(self, document_id: int, user_id: str) -> List[Dict[str, Any]]:
        """Get all chunks for a document"""
        try:
            chunks = database_service.get_document_chunks(document_id)

            # Filter by user_id for security
            user_chunks = [chunk for chunk in chunks if str(chunk['user_id']) == user_id]

            return user_chunks

        except Exception as e:
            self.logger.error(f"Failed to get document chunks: {str(e)}")
            return []

    async def regenerate_embeddings(self, user_id: str) -> Dict[str, Any]:
        """Regenerate embeddings for all user documents"""
        try:
            # Get all user documents
            documents = database_service.get_user_documents(int(user_id))

            total_regenerated = 0
            failed_documents = []

            for doc in documents:
                try:
                    # Get document chunks
                    chunks = database_service.get_document_chunks(doc['id'])

                    # Prepare for vector service
                    vector_chunks = []
                    for chunk in chunks:
                        vector_chunk = {
                            'chunk_id': chunk['id'],
                            'content': chunk['content'],
                            'document_id': chunk['document_id'],
                            'metadata': chunk['metadata']
                        }
                        vector_chunks.append(vector_chunk)

                    # Regenerate embeddings
                    if vector_chunks:
                        await vector_service.process_document_chunks(vector_chunks, user_id)
                        total_regenerated += len(vector_chunks)

                        # Update chunk embedding status
                        for chunk in chunks:
                            database_service.update_chunk_embedding_status(chunk['id'], 'completed')

                except Exception as e:
                    self.logger.error(f"Failed to regenerate embeddings for document {doc['id']}: {str(e)}")
                    failed_documents.append(doc['filename'])

            return {
                'total_regenerated': total_regenerated,
                'failed_documents': failed_documents,
                'status': 'completed' if not failed_documents else 'partial'
            }

        except Exception as e:
            self.logger.error(f"Embedding regeneration failed: {str(e)}")
            raise RAGProcessingError(f"Embedding regeneration failed: {str(e)}")

    def health_check(self) -> Dict[str, Any]:
        """Health check for RAG service"""
        try:
            return {
                'status': 'healthy',
                'document_processor': 'ok',
                'text_chunker': 'ok',
                'timestamp': datetime.now().isoformat()
            }
        except Exception as e:
            return {
                'status': 'unhealthy',
                'error': str(e),
                'timestamp': datetime.now().isoformat()
            }


# Global RAG service instance
rag_service = RAGService()